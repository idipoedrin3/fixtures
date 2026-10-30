import csv
import io
from decimal import Decimal, InvalidOperation

import openpyxl
import pandas as pd
from accounts.models import Account
from django.contrib.contenttypes.models import ContentType
from django.db import transaction
from django.db.models import Exists, OuterRef
from django.utils import timezone
from institution.models import Branch
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from taxation.models import Tax

from audit.models import AuditLog
from items.models import (
    InventoryItem,
    InventoryItemStockLayer,
    InventoryItemUnitOfMeasure,
    ItemCategory,
    ItemUnitOfMeasure,
)
from items.uom_sync import sync_inventory_item_uom_conversion_factors
from utilities.helpers import get_institution_timezone

REQUIRED_ITEM_COLS = {
    "reference_code",
    "item_name",
    "item_type",
    "item_category",
    "can_be_sold",
    "can_be_purchased",
}
REQUIRED_UOM_COLS = {"reference_code", "unit_name", "conversion_factor"}
REQUIRED_STOCK_COLS = {"reference_code", "unit_name", "quantity", "unit_cost"}
REQUIRED_BRANCH_COLS = {"reference_code", "selling_branch"}

VALID_ITEM_TYPES = {"GOODS", "SERVICES"}
BOOLEAN_CHOICES = ("True", "False")
TEMPLATE_DATA_ROWS = 500

ITEM_COLUMNS = [
    "reference_code",
    "item_name",
    "item_type",
    "item_category",
    "reorder_point",
    "inventory_account",
    "revenue_account",
    "expense_account",
    "can_be_sold",
    "sales_tax",
    "can_be_purchased",
    "purchase_tax",
]
UOM_COLUMNS = [
    "reference_code",
    "unit_name",
    "conversion_factor",
    "selling_price",
    "is_base_unit",
]
STOCK_COLUMNS = [
    "reference_code",
    "unit_name",
    "quantity",
    "unit_cost",
    "received_date",
    "expiry_date",
]
BRANCH_COLUMNS = ["reference_code", "selling_branch"]


def _str(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return None if s.lower() in ("nan", "") else s


def _decimal(value, default=Decimal("0")) -> Decimal:
    try:
        decimal_value = Decimal(str(value))

        if not decimal_value.is_finite():
            return default

        return decimal_value.quantize(Decimal("0.00000001"))
    except (InvalidOperation, TypeError, ValueError):
        return default


def _int(value, default=0) -> int:
    try:
        return int(float(str(value)))
    except (ValueError, TypeError):
        return default


def _strict_bool(value, field_name: str, default: bool) -> bool:
    if isinstance(value, bool):
        return value
    s = _str(value)
    if s is None:
        return default
    if s.lower() in ("true", "1", "yes", "y"):
        return True
    if s.lower() in ("false", "0", "no", "n"):
        return False
    raise ValueError(f"{field_name} must be either True or False.")


def _datetime(value, institution=None):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    try:
        dt = pd.to_datetime(value).to_pydatetime()
        tz = get_institution_timezone(institution)
        if timezone.is_naive(dt):
            return timezone.make_aware(dt, tz)
        return dt.astimezone(tz)
    except Exception:
        return None


def _choice_label(reference_code, name) -> str:
    reference_code = _str(reference_code)
    name = _str(name) or ""
    return f"{reference_code} | {name}" if reference_code else name


def _leaf_accounts(institution, group_type: str) -> list[Account]:
    active_children = Account.objects.filter(
        parent_id=OuterRef("pk"),
        deleted_at__isnull=True,
    )
    return list(
        Account.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
            account_role__role_group__group_type_name=group_type,
        )
        .annotate(has_children=Exists(active_children))
        .filter(has_children=False)
        .order_by("account_name", "reference_code")
    )


def _add_dropdown(
    worksheet,
    options_worksheet,
    column_index: int,
    options: list[str],
    required: bool = False,
):
    if not options:
        return

    option_column = options_worksheet.max_column + 1
    if (
        options_worksheet.max_column == 1
        and options_worksheet.max_row == 1
        and options_worksheet.cell(1, 1).value is None
    ):
        option_column = 1

    for row_index, option in enumerate(options, start=1):
        options_worksheet.cell(row=row_index, column=option_column, value=option)

    option_letter = get_column_letter(option_column)
    data_letter = get_column_letter(column_index)
    range_name = f"InventoryDropdown_{option_column}"
    worksheet.parent.defined_names.add(
        DefinedName(
            range_name,
            attr_text=(
                f"'Dropdown Options'!${option_letter}$1:"
                f"${option_letter}${len(options)}"
            ),
        )
    )
    validation = DataValidation(
        type="list",
        formula1=f"={range_name}",
        showDropDown=False,
        allow_blank=not required,
        errorStyle="stop",
        showErrorMessage=True,
        errorTitle="Invalid value",
        error="Select a value from the dropdown list.",
    )
    worksheet.add_data_validation(validation)
    validation.add(f"{data_letter}2:{data_letter}{TEMPLATE_DATA_ROWS + 1}")


def _style_sheet(worksheet, columns: list[str]):
    header_fill = PatternFill("solid", fgColor="2C3E50")
    header_font = Font(color="FFFFFF", bold=True)

    for column_index, column_name in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(
            len(column_name) + 4, 18
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def generate_bulk_inventory_template(
    format_type: str = "excel",
    institution=None,
) -> tuple[bytes, str, str]:
    categories = []
    if institution is not None:
        categories = list(
            ItemCategory.objects.filter(
                institution=institution,
                deleted_at__isnull=True,
            )
            .order_by("name")
            .values_list("name", flat=True)
        )

    sample_items = [
        {
            "reference_code": "ITM-001",
            "item_name": "Sample Item",
            "item_type": "GOODS",
            "item_category": categories[0] if categories else "",
            "reorder_point": 20,
            "inventory_account": "",
            "revenue_account": "",
            "expense_account": "",
            "can_be_sold": True,
            "sales_tax": "",
            "can_be_purchased": True,
            "purchase_tax": "",
        }
    ]
    sample_uoms = [
        {
            "reference_code": "ITM-001",
            "unit_name": "Piece",
            "conversion_factor": 1,
            "selling_price": 5000,
            "is_base_unit": True,
        },
        {
            "reference_code": "ITM-001",
            "unit_name": "Box",
            "conversion_factor": 12,
            "selling_price": 55000,
            "is_base_unit": False,
        },
    ]
    sample_stock = [
        {
            "reference_code": "ITM-001",
            "unit_name": "Piece",
            "quantity": 100,
            "unit_cost": 4000,
            "received_date": "2024-01-01",
            "expiry_date": "",
        },
        {
            "reference_code": "ITM-001",
            "unit_name": "Box",
            "quantity": 5,
            "unit_cost": 44000,
            "received_date": "2024-01-01",
            "expiry_date": "",
        },
    ]
    sample_branches = [
        {
            "reference_code": "ITM-001",
            "selling_branch": "",
        }
    ]

    if format_type == "csv":
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=ITEM_COLUMNS)
        writer.writeheader()
        writer.writerows(sample_items)
        return buf.getvalue().encode("utf-8"), "text/csv", "csv"

    buf = io.BytesIO()
    workbook = openpyxl.Workbook()
    items_sheet = workbook.active
    items_sheet.title = "Items"
    uom_sheet = workbook.create_sheet("Units of Measure")
    stock_sheet = workbook.create_sheet("Opening Stock")
    branches_sheet = workbook.create_sheet("Selling Branches")
    options_sheet = workbook.create_sheet("Dropdown Options")

    sheet_data = [
        (items_sheet, ITEM_COLUMNS, sample_items),
        (uom_sheet, UOM_COLUMNS, sample_uoms),
        (stock_sheet, STOCK_COLUMNS, sample_stock),
        (branches_sheet, BRANCH_COLUMNS, sample_branches),
    ]
    for worksheet, columns, rows in sheet_data:
        _style_sheet(worksheet, columns)
        for row_index, row in enumerate(rows, start=2):
            for column_index, column_name in enumerate(columns, start=1):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=row.get(column_name),
                )

    item_dropdowns = {
        "item_type": sorted(VALID_ITEM_TYPES),
        "can_be_sold": list(BOOLEAN_CHOICES),
        "can_be_purchased": list(BOOLEAN_CHOICES),
    }
    branch_options = []

    if institution is not None:
        asset_accounts = _leaf_accounts(institution, "asset")
        revenue_accounts = _leaf_accounts(institution, "revenue")
        expense_accounts = _leaf_accounts(institution, "expense")
        invoice_taxes = Tax.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
            is_applicable_to_invoice_or_bill=True,
            applicability__in=["invoice_only", "INVOICE_ONLY", "BOTH"],
        ).order_by("tax_name", "reference_code")
        bill_taxes = Tax.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
            is_applicable_to_invoice_or_bill=True,
            applicability__in=["bill_only", "BILL_ONLY", "BOTH"],
        ).order_by("tax_name", "reference_code")
        branches = Branch.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
        ).order_by("branch_name", "reference_code")

        item_dropdowns.update(
            {
                "item_category": categories,
                "inventory_account": [
                    _choice_label(account.reference_code, account.account_name)
                    for account in asset_accounts
                ],
                "revenue_account": [
                    _choice_label(account.reference_code, account.account_name)
                    for account in revenue_accounts
                ],
                "expense_account": [
                    _choice_label(account.reference_code, account.account_name)
                    for account in expense_accounts
                ],
                "sales_tax": [
                    _choice_label(tax.reference_code, tax.tax_name)
                    for tax in invoice_taxes
                ],
                "purchase_tax": [
                    _choice_label(tax.reference_code, tax.tax_name)
                    for tax in bill_taxes
                ],
            }
        )
        branch_options = [
            _choice_label(branch.reference_code, branch.branch_name)
            for branch in branches
        ]

    for column_name, options in item_dropdowns.items():
        _add_dropdown(
            items_sheet,
            options_sheet,
            ITEM_COLUMNS.index(column_name) + 1,
            options,
            required=column_name
            in {
                "item_type",
                "item_category",
                "can_be_sold",
                "can_be_purchased",
            },
        )

    _add_dropdown(
        uom_sheet,
        options_sheet,
        UOM_COLUMNS.index("is_base_unit") + 1,
        list(BOOLEAN_CHOICES),
    )
    _add_dropdown(
        branches_sheet,
        options_sheet,
        BRANCH_COLUMNS.index("selling_branch") + 1,
        branch_options,
    )

    options_sheet.sheet_state = "veryHidden"
    workbook.save(buf)
    buf.seek(0)
    return (
        buf.read(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )


def _parse_upload(file_bytes: bytes, filename: str):
    name = filename.lower()
    if name.endswith(".csv"):
        items_df = pd.read_csv(io.StringIO(file_bytes.decode("utf-8")))
        return items_df, None, None, None
    if name.endswith((".xlsx", ".xls")):
        xl = pd.ExcelFile(io.BytesIO(file_bytes))
        visible_sheets = [
            sheet for sheet in xl.sheet_names if sheet != "Dropdown Options"
        ]
        items_df = (
            xl.parse("Items")
            if "Items" in xl.sheet_names
            else xl.parse(visible_sheets[0])
        )
        uoms_df = (
            xl.parse("Units of Measure")
            if "Units of Measure" in xl.sheet_names
            else None
        )
        stock_df = (
            xl.parse("Opening Stock") if "Opening Stock" in xl.sheet_names else None
        )
        branches_df = (
            xl.parse("Selling Branches")
            if "Selling Branches" in xl.sheet_names
            else None
        )
        return items_df, uoms_df, stock_df, branches_df
    raise ValueError("Unsupported file format. Use .csv, .xlsx, or .xls.")


def _normalise_cols(df):
    if df is None:
        return None
    df = df.copy()
    df.columns = [str(c).strip().lower() for c in df.columns]
    return df


def _missing_cols(df, required: set) -> list:
    return sorted(required - set(df.columns))


def _object_choice_map(objects, name_attribute: str) -> dict[str, object]:
    result = {}
    names = {}

    for obj in objects:
        reference_code = _str(getattr(obj, "reference_code", None))
        name = _str(getattr(obj, name_attribute, None))
        if reference_code:
            result[reference_code.lower()] = obj
        if reference_code and name:
            result[_choice_label(reference_code, name).lower()] = obj
        if name:
            names.setdefault(name.lower(), []).append(obj)
            # Accept values from templates generated before missing codes were omitted.
            if not reference_code:
                result[f"none | {name}".lower()] = obj

    for name, matching_objects in names.items():
        if len(matching_objects) == 1:
            result[name] = matching_objects[0]

    return result


def _resolve_choice(value, choices: dict[str, object], field_name: str):
    selected = _str(value)
    if selected is None:
        return None
    obj = choices.get(selected.lower())
    if obj is None:
        raise ValueError(
            f"{field_name} '{selected}' is not a valid institution choice."
        )
    return obj


def process_bulk_inventory_upload(
    file_bytes: bytes,
    filename: str,
    request,
    dry_run: bool = False,
) -> dict:

    institution = request.institution if hasattr(request, "institution") else None

    try:
        items_df, uoms_df, stock_df, branches_df = _parse_upload(file_bytes, filename)
    except ValueError as exc:
        return {"status": "failed", "message": str(exc)}
    except Exception as exc:
        return {"status": "failed", "message": f"Error reading file: {exc}"}

    items_df = _normalise_cols(items_df)
    uoms_df = _normalise_cols(uoms_df)
    stock_df = _normalise_cols(stock_df)
    branches_df = _normalise_cols(branches_df)

    if items_df is None or items_df.empty:
        return {"status": "failed", "message": "The Items sheet is empty."}

    missing = _missing_cols(items_df, REQUIRED_ITEM_COLS)
    if missing:
        return {
            "status": "validation_failed",
            "message": f"Items sheet missing required columns: {missing}",
            "errors": [],
        }

    if uoms_df is not None and not uoms_df.empty:
        missing = _missing_cols(uoms_df, REQUIRED_UOM_COLS)
        if missing:
            return {
                "status": "validation_failed",
                "message": f"Units of Measure sheet missing required columns: {missing}",
                "errors": [],
            }

    if stock_df is not None and not stock_df.empty:
        missing = _missing_cols(stock_df, REQUIRED_STOCK_COLS)
        if missing:
            return {
                "status": "validation_failed",
                "message": f"Opening Stock sheet missing required columns: {missing}",
                "errors": [],
            }

    if branches_df is not None and not branches_df.empty:
        missing = _missing_cols(branches_df, REQUIRED_BRANCH_COLS)
        if missing:
            return {
                "status": "validation_failed",
                "message": f"Selling Branches sheet missing required columns: {missing}",
                "errors": [],
            }

    uom_rows_by_item: dict[str, list[dict]] = {}
    uom_errors: list[dict] = []
    seen_uom_rows: dict[tuple[str, str], int] = {}

    if uoms_df is not None and not uoms_df.empty:
        for idx, row in uoms_df.iterrows():
            row_num = idx + 2
            reference_code = _str(row.get("reference_code"))
            unit_name = _str(row.get("unit_name"))

            if not reference_code:
                uom_errors.append(
                    {
                        "sheet": "Units of Measure",
                        "row": row_num,
                        "error": "reference_code is required.",
                    }
                )
                continue
            if not unit_name:
                uom_errors.append(
                    {
                        "sheet": "Units of Measure",
                        "row": row_num,
                        "error": "unit_name is required.",
                    }
                )
                continue

            uom_key = (
                reference_code.casefold(),
                unit_name.casefold(),
            )
            first_row = seen_uom_rows.get(uom_key)

            if first_row is not None:
                uom_errors.append(
                    {
                        "sheet": "Units of Measure",
                        "row": row_num,
                        "error": (
                            f"Duplicate unit_name '{unit_name}' for item "
                            f"'{reference_code}'. It was first declared on "
                            f"row {first_row}. Each item can use a unit of "
                            "measure only once."
                        ),
                    }
                )
                continue

            seen_uom_rows[uom_key] = row_num

            cf = _decimal(row.get("conversion_factor", 1), default=Decimal("1"))
            if cf <= 0:
                uom_errors.append(
                    {
                        "sheet": "Units of Measure",
                        "row": row_num,
                        "error": "conversion_factor must be greater than 0.",
                    }
                )
                continue

            try:
                is_base_unit = _strict_bool(
                    row.get("is_base_unit"),
                    "is_base_unit",
                    default=False,
                )
            except ValueError as exc:
                uom_errors.append(
                    {
                        "sheet": "Units of Measure",
                        "row": row_num,
                        "error": str(exc),
                    }
                )
                continue

            uom_rows_by_item.setdefault(reference_code, []).append(
                {
                    "unit_name": unit_name,
                    "conversion_factor": cf,
                    "selling_price": _decimal(
                        row.get("selling_price", 0), default=Decimal("0")
                    ),
                    "is_base_unit": is_base_unit,
                }
            )

    if uom_errors:
        return {
            "status": "validation_failed",
            "message": "Errors found in Units of Measure sheet.",
            "errors": uom_errors,
        }

    stock_rows_by_item: dict[str, list[dict]] = {}
    stock_errors: list[dict] = []

    if stock_df is not None and not stock_df.empty:
        for idx, row in stock_df.iterrows():
            row_num = idx + 2
            reference_code = _str(row.get("reference_code"))
            unit_name = _str(row.get("unit_name"))

            if not reference_code:
                stock_errors.append(
                    {
                        "sheet": "Opening Stock",
                        "row": row_num,
                        "error": "reference_code is required.",
                    }
                )
                continue
            if not unit_name:
                stock_errors.append(
                    {
                        "sheet": "Opening Stock",
                        "row": row_num,
                        "error": "unit_name is required.",
                    }
                )
                continue

            quantity = _int(row.get("quantity", 0))
            unit_cost = _decimal(row.get("unit_cost", 0), default=Decimal("0"))

            if quantity < 0:
                stock_errors.append(
                    {
                        "sheet": "Opening Stock",
                        "row": row_num,
                        "error": "quantity cannot be negative.",
                    }
                )
                continue
            if unit_cost <= 0:
                stock_errors.append(
                    {
                        "sheet": "Opening Stock",
                        "row": row_num,
                        "error": "unit_cost must be greater than zero. "
                        "For gifted stock, enter its assessed unit value.",
                    }
                )
                continue

            received_date = (
                _datetime(row.get("received_date"), institution) or timezone.now()
            )
            if received_date > timezone.now():
                stock_errors.append(
                    {
                        "sheet": "Opening Stock",
                        "row": row_num,
                        "error": "received_date cannot be in the future.",
                    }
                )
                continue

            stock_rows_by_item.setdefault(reference_code, []).append(
                {
                    "unit_name": unit_name,
                    "quantity": quantity,
                    "unit_cost": unit_cost,
                    "received_date": received_date,
                    "expiry_date": _datetime(row.get("expiry_date"), institution),
                }
            )

    if stock_errors:
        return {
            "status": "validation_failed",
            "message": "Errors found in Opening Stock sheet.",
            "errors": stock_errors,
        }

    cross_errors = []
    for reference_code, s_rows in stock_rows_by_item.items():
        declared = {
            u["unit_name"].lower() for u in uom_rows_by_item.get(reference_code, [])
        }
        if not declared:
            continue
        for s_row in s_rows:
            if s_row["unit_name"].lower() not in declared:
                cross_errors.append(
                    {
                        "reference_code": reference_code,
                        "error": (
                            f"Opening stock references unit '{s_row['unit_name']}' which is "
                            f"not declared in the Units of Measure sheet for this item."
                        ),
                    }
                )
    if cross_errors:
        return {
            "status": "validation_failed",
            "message": "Opening Stock references undeclared units of measure.",
            "errors": cross_errors,
        }

    categories = list(
        ItemCategory.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
        )
    )
    category_choices = {
        category.name.strip().lower(): category for category in categories
    }

    asset_accounts = _leaf_accounts(institution, "asset")
    revenue_accounts = _leaf_accounts(institution, "revenue")
    expense_accounts = _leaf_accounts(institution, "expense")
    inventory_account_choices = _object_choice_map(asset_accounts, "account_name")
    revenue_account_choices = _object_choice_map(revenue_accounts, "account_name")
    expense_account_choices = _object_choice_map(expense_accounts, "account_name")

    invoice_taxes = list(
        Tax.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
            is_applicable_to_invoice_or_bill=True,
            applicability__in=["invoice_only", "INVOICE_ONLY", "BOTH"],
        )
    )
    bill_taxes = list(
        Tax.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
            is_applicable_to_invoice_or_bill=True,
            applicability__in=["bill_only", "BILL_ONLY", "BOTH"],
        )
    )
    sales_tax_choices = _object_choice_map(invoice_taxes, "tax_name")
    purchase_tax_choices = _object_choice_map(bill_taxes, "tax_name")

    institution_branches = list(
        Branch.objects.filter(
            institution=institution,
            deleted_at__isnull=True,
        )
    )
    branch_choices = _object_choice_map(institution_branches, "branch_name")
    branch_rows_by_item: dict[str, list[Branch]] = {}
    branch_errors: list[dict] = []

    if branches_df is not None and not branches_df.empty:
        for idx, row in branches_df.iterrows():
            row_num = idx + 2
            reference_code = _str(row.get("reference_code"))
            branch_value = _str(row.get("selling_branch"))

            if not reference_code and not branch_value:
                continue
            if not reference_code:
                branch_errors.append(
                    {
                        "sheet": "Selling Branches",
                        "row": row_num,
                        "error": "reference_code is required when a branch is selected.",
                    }
                )
                continue
            if not branch_value:
                continue

            try:
                branch = _resolve_choice(branch_value, branch_choices, "selling_branch")
            except ValueError as exc:
                branch_errors.append(
                    {
                        "sheet": "Selling Branches",
                        "row": row_num,
                        "error": str(exc),
                    }
                )
                continue

            existing = branch_rows_by_item.setdefault(reference_code, [])
            if branch not in existing:
                existing.append(branch)

    uploaded_reference_codes = {
        _str(value) for value in items_df["reference_code"].tolist() if _str(value)
    }
    for reference_code in branch_rows_by_item:
        if reference_code not in uploaded_reference_codes:
            branch_errors.append(
                {
                    "sheet": "Selling Branches",
                    "row": "?",
                    "error": (
                        f"reference_code '{reference_code}' does not exist "
                        "in the Items sheet."
                    ),
                }
            )

    if branch_errors:
        return {
            "status": "validation_failed",
            "message": "Errors found in Selling Branches sheet.",
            "errors": branch_errors,
        }

    existing_codes = set(
        InventoryItem.objects.filter(
            item_category__institution=institution,
            deleted_at__isnull=True,
        ).values_list("reference_code", flat=True)
    )
    existing_names = set(
        InventoryItem.objects.filter(
            item_category__institution=institution,
            deleted_at__isnull=True,
        ).values_list("item_name", flat=True)
    )

    validated_items: list[dict] = []
    item_errors: list[dict] = []
    seen_codes: set[str] = set()
    seen_names: set[str] = set()

    for idx, row in items_df.iterrows():
        row_num = idx + 2
        row_errors = []

        reference_code = _str(row.get("reference_code"))
        item_name = _str(row.get("item_name"))
        item_category_name = _str(row.get("item_category"))
        item_type = (_str(row.get("item_type")) or "GOODS").upper()
        reorder_point = _int(row.get("reorder_point"), default=0)
        try:
            can_be_sold = _strict_bool(
                row.get("can_be_sold"), "can_be_sold", default=True
            )
            can_be_purchased = _strict_bool(
                row.get("can_be_purchased"), "can_be_purchased", default=True
            )
        except ValueError as exc:
            row_errors.append(str(exc))
            can_be_sold = True
            can_be_purchased = True

        item_category = None
        inventory_account = None
        revenue_account = None
        expense_account = None
        sales_tax = None
        purchase_tax = None

        if item_category_name:
            item_category = category_choices.get(item_category_name.lower())
            if item_category is None:
                row_errors.append(
                    f"item_category '{item_category_name}' is not a valid institution choice."
                )

        choice_fields = [
            (
                "inventory_account",
                inventory_account_choices,
                row.get("inventory_account"),
            ),
            ("revenue_account", revenue_account_choices, row.get("revenue_account")),
            ("expense_account", expense_account_choices, row.get("expense_account")),
            ("sales_tax", sales_tax_choices, row.get("sales_tax")),
            ("purchase_tax", purchase_tax_choices, row.get("purchase_tax")),
        ]
        resolved_choices = {}
        for field_name, choices, value in choice_fields:
            try:
                resolved_choices[field_name] = _resolve_choice(
                    value, choices, field_name
                )
            except ValueError as exc:
                row_errors.append(str(exc))
                resolved_choices[field_name] = None

        inventory_account = resolved_choices["inventory_account"]
        revenue_account = resolved_choices["revenue_account"]
        expense_account = resolved_choices["expense_account"]
        sales_tax = resolved_choices["sales_tax"]
        purchase_tax = resolved_choices["purchase_tax"]

        if not reference_code:
            row_errors.append("reference_code is required.")
        elif reference_code in seen_codes:
            row_errors.append(
                f"Duplicate reference_code '{reference_code}' within the file."
            )
        elif reference_code in existing_codes:
            row_errors.append(
                f"reference_code '{reference_code}' already exists in the system."
            )

        if not item_name:
            row_errors.append("item_name is required.")
        elif item_name in seen_names:
            row_errors.append(f"Duplicate item_name '{item_name}' within the file.")
        elif item_name in existing_names:
            row_errors.append(f"item_name '{item_name}' already exists in the system.")

        if not item_category_name:
            row_errors.append("item_category is required.")
        if item_type not in VALID_ITEM_TYPES:
            row_errors.append(f"item_type must be one of {sorted(VALID_ITEM_TYPES)}.")
        if item_type == "SERVICES" and inventory_account is not None:
            row_errors.append("inventory_account cannot be set for SERVICES items.")
        if not can_be_sold and sales_tax is not None:
            row_errors.append("sales_tax cannot be set when can_be_sold is False.")
        if not can_be_purchased and purchase_tax is not None:
            row_errors.append(
                "purchase_tax cannot be set when can_be_purchased is False."
            )

        uom_rows = uom_rows_by_item.get(reference_code, [])
        if uom_rows:
            base_count = sum(1 for u in uom_rows if u["is_base_unit"])
            if base_count == 0:
                row_errors.append(
                    "At least one unit of measure must be marked is_base_unit = True."
                )
            elif base_count > 1:
                row_errors.append(
                    "Only one unit of measure can be set as the base unit."
                )
            else:
                base_row = next(u for u in uom_rows if u["is_base_unit"])
                if base_row["conversion_factor"] != Decimal("1"):
                    row_errors.append("The base unit conversion_factor must be 1.")

        if item_type == "SERVICES" and stock_rows_by_item.get(reference_code):
            row_errors.append("Opening stock cannot be added for SERVICES items.")

        if row_errors:
            item_errors.append(
                {"row": row_num, "reference_code": reference_code, "errors": row_errors}
            )
            continue

        if reference_code:
            seen_codes.add(reference_code)
        if item_name:
            seen_names.add(item_name)

        validated_items.append(
            {
                "reference_code": reference_code,
                "item_name": item_name,
                "item_type": item_type,
                "item_category": item_category,
                "reorder_point": reorder_point,
                "inventory_account": inventory_account,
                "revenue_account": revenue_account,
                "expense_account": expense_account,
                "can_be_sold": can_be_sold,
                "sales_tax": sales_tax,
                "can_be_purchased": can_be_purchased,
                "purchase_tax": purchase_tax,
                "selling_branches": branch_rows_by_item.get(reference_code),
                "uom_rows": uom_rows,
                "stock_rows": stock_rows_by_item.get(reference_code, []),
            }
        )

    if item_errors:
        return {
            "status": "validation_failed",
            "message": f"{len(item_errors)} row(s) failed validation.",
            "errors": item_errors,
        }

    if not validated_items:
        return {"status": "failed", "message": "No valid items found to create."}

    if dry_run:
        return {
            "status": "success",
            "message": f"{len(validated_items)} item(s) passed validation.",
        }

    total_stock_layers = 0

    try:
        with transaction.atomic():
            all_uom_names = {
                u["unit_name"] for d in validated_items for u in d["uom_rows"]
            } | {s["unit_name"] for d in validated_items for s in d["stock_rows"]}
            uom_cache: dict[str, ItemUnitOfMeasure] = {}
            for unit_name in all_uom_names:
                obj = ItemUnitOfMeasure.objects.filter(
                    unit_name__iexact=unit_name,
                    institution=institution,
                    deleted_at__isnull=True,
                ).first()
                if not obj:
                    obj = ItemUnitOfMeasure.objects.create(
                        unit_name=unit_name,
                        unit_abbreviation=unit_name[:10],
                        institution=institution,
                    )
                    obj.confirm_create()
                uom_cache[unit_name.lower()] = obj

            created_items = []
            content_type = ContentType.objects.get_for_model(InventoryItem)
            audit_logs = []

            for d in validated_items:

                item_obj = InventoryItem.objects.create(
                    reference_code=d["reference_code"],
                    item_name=d["item_name"],
                    item_type=d["item_type"],
                    item_category=d["item_category"],
                    reorder_point=d["reorder_point"],
                    inventory_account=d["inventory_account"],
                    revenue_account=d["revenue_account"],
                    expense_account=d["expense_account"],
                    can_be_sold=d["can_be_sold"],
                    sales_tax=d["sales_tax"],
                    can_be_purchased=d["can_be_purchased"],
                    purchase_tax=d["purchase_tax"],
                    institution=institution,
                    is_active=True,
                    approval_status="active",
                )

                created_items.append(item_obj)

                item_obj.assign_inventory_items_to_all_branches(
                    branches=d["selling_branches"],
                    skip_uom_sync=True,
                )

                if d["uom_rows"]:
                    base_uom_row = next(
                        row for row in d["uom_rows"] if row["is_base_unit"]
                    )
                    base_uom = uom_cache[base_uom_row["unit_name"].lower()]

                    units_to_create = [
                        InventoryItemUnitOfMeasure(
                            inventory_item=item_obj,
                            unit_of_measure=uom_cache[uom_row["unit_name"].lower()],
                            conversion_factor=(
                                Decimal("1")
                                if uom_row["is_base_unit"]
                                else Decimal("0")
                            ),
                            selling_price=uom_row["selling_price"],
                            is_base_unit_of_measure=uom_row["is_base_unit"],
                            other_quantity=(
                                None
                                if uom_row["is_base_unit"]
                                else uom_row["conversion_factor"]
                            ),
                            other_uom=(None if uom_row["is_base_unit"] else base_uom),
                        )
                        for uom_row in d["uom_rows"]
                    ]
                    created_units = InventoryItemUnitOfMeasure.objects.bulk_create(
                        units_to_create
                    )

                    sync_inventory_item_uom_conversion_factors(item_obj.id)

                    for unit in created_units:
                        unit.refresh_from_db(
                            fields=[
                                "conversion_factor",
                                "other_quantity",
                                "other_uom",
                            ]
                        )
                        unit.sync_to_branches()

                for s_row in d["stock_rows"]:
                    uom_obj = uom_cache.get(s_row["unit_name"].lower())
                    quantity = s_row["quantity"]
                    unit_cost = s_row["unit_cost"]

                    existing_layer = (
                        InventoryItemStockLayer.objects.select_for_update()
                        .filter(
                            inventory_item=item_obj,
                            unit_of_measure=uom_obj,
                            unit_cost=unit_cost,
                            deleted_at__isnull=True,
                        )
                        .first()
                    )
                    if existing_layer:
                        existing_layer.quantity_remaining += quantity
                        existing_layer.save(update_fields=["quantity_remaining"])
                    else:
                        InventoryItemStockLayer.objects.create(
                            inventory_item=item_obj,
                            unit_of_measure=uom_obj,
                            quantity_remaining=quantity,
                            unit_cost=unit_cost,
                            recieved_date=s_row["received_date"],
                            expiry_date=s_row["expiry_date"],
                        )
                    total_stock_layers += 1

                audit_logs.append(
                    AuditLog(
                        content_type=content_type,
                        object_id=item_obj.pk,
                        action="CREATE",
                        user=request.user if request.user.is_authenticated else None,
                        description=f"Bulk created Inventory Item: {item_obj.item_name} ({item_obj.reference_code})",
                    )
                )

            AuditLog.objects.bulk_create(audit_logs)

        return {
            "status": "success",
            "message": (
                f"Successfully created {len(created_items)} inventory item(s)"
                + (
                    f" with {total_stock_layers} opening stock layer(s)."
                    if total_stock_layers
                    else "."
                )
            ),
            "created": len(created_items),
            "stock_layers_created": total_stock_layers,
        }

    except Exception as exc:
        return {"status": "failed", "message": f"Database error during creation: {exc}"}


def process_bulk_inventory_from_file(uploaded_file, request) -> dict:
    file_bytes = uploaded_file.read()
    return process_bulk_inventory_upload(file_bytes, uploaded_file.name, request)
