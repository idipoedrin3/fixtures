import base64
from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.template.loader import render_to_string
from weasyprint import HTML
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill, Protection

from reports.exporting import GeneratedReport
from .models import BranchInventoryItemUnitOfMeasure, StockMovementLog


def get_stock_movement_datetime_bounds(start_date=None, end_date=None):
    if isinstance(start_date, str):
        start_date = date.fromisoformat(start_date)

    if isinstance(end_date, str):
        end_date = date.fromisoformat(end_date)

    current_timezone = timezone.get_current_timezone()
    start_datetime = (
        timezone.make_aware(
            datetime.combine(start_date, time.min),
            current_timezone,
        )
        if start_date
        else None
    )
    end_datetime = (
        timezone.make_aware(
            datetime.combine(end_date + timedelta(days=1), time.min),
            current_timezone,
        )
        if end_date
        else None
    )

    return start_datetime, end_datetime


def build_stock_movement_queryset(
    institution,
    branch_uuid=None,
    product_uuid=None,
    movement_type=None,
    stock_level=None,
    start_date=None,
    end_date=None,
):
    logs = (
        StockMovementLog.objects.undeleted()
        .filter(institution=institution)
        .select_related(
            "inventory_item",
            "unit_of_measure",
            "branch",
            "institution",
        )
    )

    if branch_uuid:
        logs = logs.filter(
            branch__base_uuid=branch_uuid,
            branch__institution=institution,
        )
    if product_uuid:
        logs = logs.filter(inventory_item__base_uuid=product_uuid)
    if movement_type:
        logs = logs.filter(movement_type=movement_type)
    if stock_level:
        logs = logs.filter(stock_level=stock_level)

    start_datetime, end_datetime = get_stock_movement_datetime_bounds(
        start_date,
        end_date,
    )
    if start_datetime:
        logs = logs.filter(movement_date__gte=start_datetime)
    if end_datetime:
        logs = logs.filter(movement_date__lt=end_datetime)

    return logs.order_by("-movement_date", "-id")


def _institution_logo_base64(institution):
    logo = getattr(institution, "institution_logo", None)
    if not logo:
        return None

    try:
        logo.open("rb")
        encoded = base64.b64encode(logo.read()).decode("ascii")
        logo.close()
        return encoded
    except (FileNotFoundError, OSError, ValueError):
        return None


def sync_uoms_to_branch_item(branch_item, institution_uoms):
    existing_uom_ids = set(
        BranchInventoryItemUnitOfMeasure.objects.filter(
            branch_inventory_item=branch_item,
            deleted_at__isnull=True,
        ).values_list("unit_of_measure_id", flat=True)
    )

    to_create = [
        BranchInventoryItemUnitOfMeasure(
            branch_inventory_item=branch_item,
            unit_of_measure=uom.unit_of_measure,
            conversion_factor=uom.conversion_factor,
            selling_price=uom.selling_price,
            is_base_unit_of_measure=uom.is_base_unit_of_measure,
        )
        for uom in institution_uoms
        if uom.unit_of_measure_id not in existing_uom_ids
    ]

    if to_create:
        BranchInventoryItemUnitOfMeasure.objects.bulk_create(
            to_create, ignore_conflicts=True
        )


def generate_stock_movement_file(
    institution,
    logs,
    file_type="pdf",
    start_date=None,
    end_date=None,
    filters=None,
):
    log_data = []
    total_incoming = 0
    total_outgoing = 0

    for log in logs.iterator(chunk_size=500):
        quantity_change = log.quantity_change or 0
        if quantity_change > 0:
            total_incoming += quantity_change
        elif quantity_change < 0:
            total_outgoing += abs(quantity_change)

        log_data.append(
            {
                "product_name": log.inventory_item.item_name,
                "movement_type": log.get_movement_type_display(),
                "stock_level": log.get_stock_level_display(),
                "source_reference": log.source_reference,
                "quantity_change": quantity_change,
                "quantity_before": log.quantity_before,
                "quantity_after": log.quantity_after,
                "unit_cost": log.unit_cost,
                "unit_of_measure": (
                    log.unit_of_measure.unit_name if log.unit_of_measure else None
                ),
                "movement_date": timezone.localtime(log.movement_date).strftime(
                    "%Y-%m-%d %H:%M"
                ),
                "location": log.location,
            }
        )

    summary = {
        "movement_count": len(log_data),
        "total_incoming": total_incoming,
        "total_outgoing": total_outgoing,
    }

    if file_type == "excel":
        buffer = _generate_stock_movement_excel(
            institution,
            log_data,
            start_date,
            end_date,
        )
        extension = "xlsx"
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        buffer = _generate_stock_movement_pdf(
            institution,
            log_data,
            start_date,
            end_date,
            filters=filters or {},
            summary=summary,
        )

        extension = "pdf"
        content_type = "application/pdf"

    timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
    return GeneratedReport(
        buffer=buffer,
        filename=f"StockMovementReport_{timestamp}.{extension}",
        content_type=content_type,
    )


def _generate_stock_movement_pdf(
    institution,
    log_data,
    start_date=None,
    end_date=None,
    filters=None,
    summary=None,
):
    html_string = render_to_string(
        "stock_movement.html",
        {
            "institution": institution,
            "data": log_data,
            "start_date": start_date,
            "end_date": end_date,
            "filters": filters or {},
            "summary": summary or {},
            "institution_logo_base64": _institution_logo_base64(institution),
            "generated_at": timezone.localtime().strftime("%Y-%m-%d %H:%M"),
        },
    )

    return BytesIO(HTML(string=html_string).write_pdf())


def _generate_stock_movement_excel(
    institution,
    log_data,
    start_date=None,
    end_date=None,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Movement"

    row_idx = 1

    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
    cell = ws.cell(row_idx, 1, institution.institution_name)
    cell.alignment = Alignment(horizontal="center")
    cell.font = Font(bold=True, size=16, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1C2B3A")
    cell.protection = Protection(locked=True)
    ws.row_dimensions[row_idx].height = 26
    row_idx += 1

    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=11)
    cell = ws.cell(row_idx, 1, "Stock Movement Log")
    cell.alignment = Alignment(horizontal="center")
    cell.font = Font(bold=True, size=12, color="1C2B3A")
    cell.protection = Protection(locked=True)
    row_idx += 1

    if start_date and end_date:

        ws.merge_cells(
            start_row=row_idx, start_column=1, end_row=row_idx, end_column=11
        )
        cell = ws.cell(row_idx, 1, f"Date Range: {start_date} - {end_date}")
        cell.alignment = Alignment(horizontal="center")
        cell.protection = Protection(locked=True)

        row_idx += 1

    headers = [
        "Date",
        "Product",
        "Movement Type",
        "Stock Level",
        "Location",
        "Unit",
        "Quantity Change",
        "Quantity Before",
        "Quantity After",
        "Unit Cost",
        "Source Reference",
    ]
    header_row = row_idx
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row_idx, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C2B3A")
        cell.alignment = Alignment(vertical="center")
        cell.protection = Protection(locked=True)
    row_idx += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    for log in log_data:
        ws.cell(row_idx, 1, log["movement_date"])
        ws.cell(row_idx, 2, log["product_name"])
        ws.cell(row_idx, 3, log["movement_type"])
        ws.cell(row_idx, 4, log.get("stock_level", "-"))
        ws.cell(row_idx, 5, log.get("location", "-"))
        ws.cell(row_idx, 6, log.get("unit_of_measure", "-"))
        cell = ws.cell(row_idx, 7, log.get("quantity_change", 0))
        cell.number_format = "#,##0"
        cell = ws.cell(row_idx, 8, log.get("quantity_before", 0))
        cell.number_format = "#,##0"
        cell = ws.cell(row_idx, 9, log.get("quantity_after", 0))
        cell.number_format = "#,##0"
        cell = ws.cell(row_idx, 10, log.get("unit_cost", 0))
        cell.number_format = "#,##0.00"
        ws.cell(row_idx, 11, log.get("source_reference", "-"))
        cell.protection = Protection(locked=False)

        row_idx += 1

    column_widths = {
        1: 18,
        2: 25,
        3: 20,
        4: 18,
        5: 25,
        6: 15,
        7: 18,
        8: 18,
        9: 18,
        10: 18,
        11: 25,
    }

    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.protection.enable()

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
